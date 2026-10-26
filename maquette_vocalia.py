"""Maquette produit Vocalia — présentation PowerPoint."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

W, H = Inches(13.333), Inches(7.5)
INK = RGBColor(0x14, 0x14, 0x14)
MUTED = RGBColor(0x6B, 0x6B, 0x6B)
LINE = RGBColor(0xE4, 0xE4, 0xE4)
BG = RGBColor(0xFF, 0xFF, 0xFF)
PAPER = RGBColor(0xF7, 0xF7, 0xF6)
SIDE = RGBColor(0xF1, 0xF1, 0xF0)
ACCENT = RGBColor(0x1F, 0x6B, 0x5A)
ACCENT_D = RGBColor(0x16, 0x4B, 0x40)
SOFT = RGBColor(0xE7, 0xF2, 0xEE)
GOLD = RGBColor(0xA6, 0x84, 0x46)
DARK = RGBColor(0x0B, 0x0B, 0x0B)
DARK_CARD = RGBColor(0x16, 0x16, 0x16)
DARK_LINE = RGBColor(0x2C, 0x2C, 0x2C)
OFF = RGBColor(0xF4, 0xF4, 0xF3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SERIF = "Georgia"
SANS = "Calibri"


def fill(shape, color, stroke=None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if stroke is None:
        shape.line.fill.background()
    else:
        shape.line.fill.solid()
        shape.line.color.rgb = stroke
        shape.line.width = Pt(0.75)


def round_rect(slide, l, t, w, h, color, stroke=None, adj=0.08):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    fill(s, color, stroke)
    try:
        s.adjustments[0] = adj
    except Exception:
        pass
    return s


def rect(slide, l, t, w, h, color, stroke=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    fill(s, color, stroke)
    return s


def oval(slide, l, t, w, h, color, stroke=None):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, l, t, w, h)
    fill(s, color, stroke)
    return s


def write(shape, text, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT, font=SANS, anchor=MSO_ANCHOR.TOP):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}.get(anchor, "t"))
    except Exception:
        pass
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(2)
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold
        run.font.name = font
    return tf


def text(slide, l, t, w, h, content, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT, font=SANS, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(l, t, w, h)
    write(box, content, size, color, bold, align, font, anchor)
    return box


def footer(slide, n, total, dark=False):
    c = RGBColor(0x9A, 0x9A, 0x9A) if not dark else RGBColor(0x7A, 0x7A, 0x7A)
    text(slide, Inches(0.55), Inches(7.12), Inches(8), Inches(0.28), "Vocalia  ·  Maquette produit", 10, c, False, font=SANS)
    text(slide, Inches(10.6), Inches(7.12), Inches(2.2), Inches(0.28), f"{n:02d}  /  {total:02d}", 10, c, False, PP_ALIGN.RIGHT)


def kicker(slide, label, dark=False):
    text(slide, Inches(0.55), Inches(0.28), Inches(12), Inches(0.28), label.upper(), 11, ACCENT if not dark else GOLD, True)


def heading(slide, title, subtitle=None, dark=False):
    ink = WHITE if dark else INK
    mut = RGBColor(0xA8, 0xA8, 0xA8) if dark else MUTED
    text(slide, Inches(0.55), Inches(0.52), Inches(12.2), Inches(0.58), title, 32, ink, False, font=SERIF)
    if subtitle:
        text(slide, Inches(0.55), Inches(1.12), Inches(12.2), Inches(0.4), subtitle, 15, mut)


def chrome(slide, active, title, subtitle, dark=False):
    bg = DARK if dark else PAPER
    side = RGBColor(0x11, 0x11, 0x11) if dark else SIDE
    ink = WHITE if dark else INK
    mut = RGBColor(0x9A, 0x9A, 0x9A) if dark else MUTED
    card = DARK_CARD if dark else WHITE
    line = DARK_LINE if dark else LINE
    rect(slide, 0, 0, W, H, bg)
    round_rect(slide, Inches(0.35), Inches(0.28), Inches(12.65), Inches(6.7), card if not dark else DARK, line, 0.03)
    rect(slide, Inches(0.35), Inches(0.28), Inches(2.45), Inches(6.7), side)
    text(slide, Inches(0.55), Inches(0.48), Inches(2.05), Inches(0.32), "Vocalia", 15, WHITE if dark else INK, False, font=SERIF)
    text(slide, Inches(0.55), Inches(0.78), Inches(2.05), Inches(0.22), "Transcription", 9, mut)
    nav = ["Transcription", "Assistant", "Archives", "Réglages"]
    for i, name in enumerate(nav):
        y = Inches(1.28) + Inches(i * 0.5)
        on = name == active
        if on:
            round_rect(slide, Inches(0.48), y, Inches(2.18), Inches(0.42), ACCENT if dark else WHITE, None, 0.18)
        text(slide, Inches(0.65), y + Inches(0.07), Inches(1.9), Inches(0.3), name, 12, WHITE if dark or on else mut, on)
    text(slide, Inches(0.55), Inches(6.28), Inches(2.05), Inches(0.22), "Stéphane A.", 11, WHITE if dark else INK, True)
    text(slide, Inches(0.55), Inches(6.5), Inches(2.05), Inches(0.22), "Se déconnecter", 10, mut)
    oval(slide, Inches(12.28), Inches(0.48), Inches(0.38), Inches(0.38), OFF if not dark else DARK_CARD, line)
    text(slide, Inches(2.95), Inches(0.42), Inches(8.8), Inches(0.42), title, 22, ink, False, font=SERIF)
    text(slide, Inches(2.95), Inches(0.86), Inches(9.2), Inches(0.32), subtitle, 12, mut)
    return card, ink, mut, line


def cover(prs, n, total):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, PAPER)
    rect(s, 0, 0, Inches(0.14), H, ACCENT)
    rect(s, Inches(8.9), 0, Inches(4.44), H, INK)
    text(s, Inches(0.7), Inches(1.55), Inches(7.6), Inches(0.3), "MAQUETTE PRODUIT", 12, ACCENT, True)
    text(s, Inches(0.7), Inches(2.0), Inches(7.8), Inches(1.15), "Vocalia", 66, INK, False, font=SERIF)
    text(
        s,
        Inches(0.7),
        Inches(3.35),
        Inches(7.6),
        Inches(1.2),
        "Application de transcription pour la parole réelle :\nprédications, réunions, entretiens, vidéos.",
        18,
        MUTED,
    )
    text(s, Inches(0.7), Inches(5.15), Inches(7.6), Inches(0.9), "Web  ·  Android  ·  iOS\nCompte privé  ·  Archives  ·  Partage", 14, INK)
    text(s, Inches(9.25), Inches(2.4), Inches(3.8), Inches(0.4), "Le texte,", 22, WHITE, False, font=SERIF)
    text(s, Inches(9.25), Inches(2.85), Inches(3.8), Inches(0.5), "après la parole.", 22, GOLD, False, font=SERIF)
    text(s, Inches(9.25), Inches(4.0), Inches(3.8), Inches(1.6), "Micro, fichier ou lien.\nTitre automatique.\nWhatsApp, copie, PDF.", 15, RGBColor(0xC8, 0xC8, 0xC8))
    text(s, Inches(0.7), Inches(6.85), Inches(7.5), Inches(0.3), "Confidentiel  ·  usage interne  ·  2026", 11, MUTED)
    footer(s, n, total)


def sommaire(prs, n, total):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, PAPER)
    kicker(s, "Vue d’ensemble")
    heading(s, "Sommaire")
    items = [
        ("01", "Identité et parcours"),
        ("02", "Connexion et inscription"),
        ("03", "Studio de transcription"),
        ("04", "Titre, outils et partage"),
        ("05", "Assistant"),
        ("06", "Archives"),
        ("07", "Réglages et thème"),
        ("08", "Mobile natif"),
    ]
    for i, (num, title) in enumerate(items):
        col, row = i % 2, i // 2
        x = Inches(0.55) + Inches(col * 6.35)
        y = Inches(1.85) + Inches(row * 1.15)
        round_rect(s, x, y, Inches(6.05), Inches(0.95), WHITE, LINE, 0.08)
        text(s, x + Inches(0.28), y + Inches(0.28), Inches(0.7), Inches(0.4), num, 18, ACCENT, True)
        text(s, x + Inches(1.1), y + Inches(0.3), Inches(4.6), Inches(0.4), title, 18, INK)
    footer(s, n, total)


def parcours(prs, n, total):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, PAPER)
    kicker(s, "Expérience")
    heading(s, "Un fil unique", "De l’ouverture au partage, sans changer d’outil.")
    steps = [
        ("01", "Arrivée", "Logo, puis connexion."),
        ("02", "Capturer", "Micro, fichier ou lien."),
        ("03", "Corriger", "Le texte reste éditable."),
        ("04", "Nommer", "Titre proposé, puis archive."),
        ("05", "Diffuser", "Copie, WhatsApp, PDF."),
    ]
    for i, (num, title, body) in enumerate(steps):
        x = Inches(0.5) + Inches(i * 2.52)
        round_rect(s, x, Inches(2.05), Inches(2.35), Inches(3.55), WHITE, LINE, 0.08)
        text(s, x + Inches(0.18), Inches(2.25), Inches(2), Inches(0.3), num, 12, ACCENT, True)
        text(s, x + Inches(0.18), Inches(2.7), Inches(2), Inches(0.7), title, 20, INK, False, font=SERIF)
        text(s, x + Inches(0.18), Inches(3.55), Inches(2), Inches(1.5), body, 14, MUTED)
    footer(s, n, total)


def splash(prs, n, total):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, PAPER)
    kicker(s, "01  ·  Ouverture")
    heading(s, "Splash", "Le logo seul. Aucun halo, aucun carré d’accent.")
    round_rect(s, Inches(4.15), Inches(1.85), Inches(5.05), Inches(4.55), WHITE, LINE, 0.06)
    oval(s, Inches(6.05), Inches(2.55), Inches(1.2), Inches(1.2), WHITE, ACCENT)
    text(s, Inches(6.05), Inches(2.72), Inches(1.2), Inches(0.9), "V", 34, INK, False, PP_ALIGN.CENTER, SERIF, MSO_ANCHOR.MIDDLE)
    text(s, Inches(4.3), Inches(3.95), Inches(4.75), Inches(0.5), "Vocalia", 28, INK, False, PP_ALIGN.CENTER, SERIF)
    text(s, Inches(4.3), Inches(4.5), Inches(4.75), Inches(0.3), "TRANSCRIPTION", 11, MUTED, True, PP_ALIGN.CENTER)
    footer(s, n, total)


def auth(prs, n, total, signup=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, PAPER)
    kicker(s, "02  ·  Accès")
    heading(s, "Inscription" if signup else "Connexion", "Logo à gauche. Thème à droite. Rien sur le serveur dans le formulaire.")
    round_rect(s, Inches(0.55), Inches(1.7), Inches(12.25), Inches(5.15), WHITE, LINE, 0.04)
    rect(s, Inches(0.55), Inches(1.7), Inches(12.25), Inches(0.62), WHITE, LINE)
    text(s, Inches(0.8), Inches(1.82), Inches(4), Inches(0.38), "Vocalia", 16, INK, False, font=SERIF)
    oval(s, Inches(12.15), Inches(1.82), Inches(0.38), Inches(0.38), OFF, LINE)
    rect(s, Inches(0.55), Inches(2.32), Inches(4.7), Inches(4.53), SIDE)
    text(s, Inches(0.85), Inches(2.6), Inches(4.1), Inches(1.3), "Le texte,\naprès la parole.", 26, INK, False, font=SERIF)
    text(s, Inches(0.85), Inches(4.05), Inches(4.1), Inches(1.1), "Prédications, réunions, entretiens.\nTranscrire, résumer, archiver.", 13, MUTED)
    text(s, Inches(0.85), Inches(5.35), Inches(4.1), Inches(1.1), "Micro, fichier ou lien\nRésumé et traduction\nArchives liées au compte", 13, ACCENT_D)
    text(s, Inches(5.55), Inches(2.55), Inches(6.8), Inches(0.45), "Créer un compte" if signup else "Bon retour.", 24, INK, False, font=SERIF)
    fields = ["Prénom", "Nom", "Email", "Mot de passe"] if signup else ["Email", "Mot de passe"]
    y = Inches(3.15)
    for field in fields:
        text(s, Inches(5.55), y, Inches(6.7), Inches(0.22), field, 11, MUTED)
        round_rect(s, Inches(5.55), y + Inches(0.22), Inches(6.7), Inches(0.38), PAPER, LINE, 0.12)
        y += Inches(0.68)
    round_rect(s, Inches(5.55), y + Inches(0.05), Inches(6.7), Inches(0.44), ACCENT, None, 0.12)
    text(s, Inches(5.55), y + Inches(0.12), Inches(6.7), Inches(0.32), "Créer mon compte" if signup else "Se connecter", 13, WHITE, True, PP_ALIGN.CENTER)
    footer(s, n, total)


def studio(prs, n, total):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    card, ink, mut, line = chrome(s, "Transcription", "Transcription", "Micro, fichier ou lien. Le texte s’affiche à droite.")
    round_rect(s, Inches(2.95), Inches(1.28), Inches(3.35), Inches(5.35), PAPER, line, 0.06)
    for i, name in enumerate(["Micro", "Fichier", "Lien"]):
        x = Inches(3.1) + Inches(i * 1.05)
        on = i == 0
        round_rect(s, x, Inches(1.48), Inches(0.95), Inches(0.34), SOFT if on else WHITE, ACCENT if on else line, 0.4)
        text(s, x, Inches(1.5), Inches(0.95), Inches(0.3), name, 10, ACCENT_D if on else MUTED, True, PP_ALIGN.CENTER)
    text(s, Inches(3.1), Inches(2.15), Inches(3.05), Inches(0.4), "00:00", 22, ink, False, PP_ALIGN.CENTER, SERIF)
    text(s, Inches(3.1), Inches(2.55), Inches(3.05), Inches(0.4), "Arrêt automatique, 30 min.", 11, mut, False, PP_ALIGN.CENTER)
    round_rect(s, Inches(3.25), Inches(3.1), Inches(2.75), Inches(0.42), ACCENT, None, 0.2)
    text(s, Inches(3.25), Inches(3.16), Inches(2.75), Inches(0.32), "Commencer", 13, WHITE, True, PP_ALIGN.CENTER)
    text(s, Inches(3.15), Inches(3.75), Inches(3), Inches(0.22), "Durée maximale", 10, mut)
    round_rect(s, Inches(3.15), Inches(4.0), Inches(2.95), Inches(0.36), WHITE, line, 0.12)
    text(s, Inches(3.25), Inches(4.04), Inches(2.75), Inches(0.28), "30 minutes", 12, ink)
    text(s, Inches(3.15), Inches(4.5), Inches(3), Inches(0.22), "Langue", 10, mut)
    round_rect(s, Inches(3.15), Inches(4.75), Inches(2.95), Inches(0.36), WHITE, line, 0.12)
    text(s, Inches(3.25), Inches(4.79), Inches(2.75), Inches(0.28), "Détection automatique", 12, ink)
    round_rect(s, Inches(6.5), Inches(1.28), Inches(6.15), Inches(5.35), PAPER, line, 0.06)
    text(s, Inches(6.7), Inches(3.2), Inches(5.75), Inches(0.4), "Rien pour l’instant.", 18, ink, False, PP_ALIGN.CENTER, SERIF)
    text(s, Inches(6.9), Inches(3.7), Inches(5.35), Inches(0.8), "Le texte transcrit apparaît ici.\nCorrigeable avant archivage.", 13, mut, False, PP_ALIGN.CENTER)
    footer(s, n, total)


def document(prs, n, total):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    card, ink, mut, line = chrome(s, "Transcription", "Transcription", "Titre proposé automatiquement. Partage sans quitter l’écran.")
    round_rect(s, Inches(2.95), Inches(1.28), Inches(9.7), Inches(5.35), PAPER, line, 0.05)
    text(s, Inches(3.15), Inches(1.42), Inches(9.3), Inches(0.25), "248 mots   ·   français", 11, mut)
    text(s, Inches(3.15), Inches(1.72), Inches(9.3), Inches(0.22), "Titre", 10, mut)
    round_rect(s, Inches(3.15), Inches(1.95), Inches(9.3), Inches(0.4), WHITE, line, 0.1)
    text(s, Inches(3.3), Inches(2.0), Inches(9), Inches(0.3), "Accueil du dimanche — persévérance", 13, ink)
    round_rect(s, Inches(3.15), Inches(2.5), Inches(9.3), Inches(2.15), WHITE, line, 0.06)
    text(
        s,
        Inches(3.35),
        Inches(2.65),
        Inches(8.95),
        Inches(1.9),
        "Frères et sœurs, ce matin nous parlons de la persévérance.\n"
        "La parole n’est pas seulement entendue : elle est gardée, relue, partagée.\n"
        "Trois mouvements : rester, prier, servir ensemble.",
        13,
        ink,
    )
    actions = [("Copier", False), ("WhatsApp", False), ("PDF", False), ("Lire", False), ("Archiver", True)]
    for i, (name, primary) in enumerate(actions):
        x = Inches(3.15) + Inches(i * 1.85)
        round_rect(s, x, Inches(4.85), Inches(1.7), Inches(0.38), ACCENT if primary else WHITE, None if primary else line, 0.25)
        text(s, x, Inches(4.9), Inches(1.7), Inches(0.3), name, 11, WHITE if primary else ink, True, PP_ALIGN.CENTER)
    text(s, Inches(3.15), Inches(5.4), Inches(2.2), Inches(0.22), "Langue du résultat", 10, mut)
    round_rect(s, Inches(3.15), Inches(5.65), Inches(2.2), Inches(0.34), WHITE, line, 0.12)
    text(s, Inches(3.25), Inches(5.68), Inches(2), Inches(0.28), "français", 12, ink)
    round_rect(s, Inches(5.5), Inches(5.65), Inches(1.45), Inches(0.34), WHITE, line, 0.2)
    text(s, Inches(5.5), Inches(5.68), Inches(1.45), Inches(0.28), "Résumer", 12, ink, True, PP_ALIGN.CENTER)
    round_rect(s, Inches(7.1), Inches(5.65), Inches(1.45), Inches(0.34), WHITE, line, 0.2)
    text(s, Inches(7.1), Inches(5.68), Inches(1.45), Inches(0.28), "Traduire", 12, ink, True, PP_ALIGN.CENTER)
    footer(s, n, total)


def assistant(prs, n, total):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    card, ink, mut, line = chrome(s, "Assistant", "Assistant", "Question écrite ou vocale. Réponse simple, lisible à voix haute.")
    round_rect(s, Inches(5.35), Inches(1.4), Inches(7.15), Inches(0.85), SOFT, None, 0.1)
    text(s, Inches(5.55), Inches(1.52), Inches(6.8), Inches(0.6), "Vous  —  Quels points retenir d’une prédication sur la persévérance ?", 13, ink)
    round_rect(s, Inches(3.05), Inches(2.45), Inches(7.4), Inches(1.55), PAPER, line, 0.1)
    text(s, Inches(3.25), Inches(2.58), Inches(7.05), Inches(1.3), "Vocalia  —  Trois points : rester malgré la fatigue, prier avec des mots simples, servir ensemble. Titre possible : « Tenir, même quand c’est long ».", 13, ink)
    text(s, Inches(3.05), Inches(4.2), Inches(9.4), Inches(0.3), "Copier     WhatsApp     PDF     Archiver     Nouvelle conversation", 12, mut)
    round_rect(s, Inches(2.95), Inches(5.55), Inches(9.7), Inches(0.85), PAPER, line, 0.18)
    oval(s, Inches(3.15), Inches(5.7), Inches(0.52), Inches(0.52), WHITE, line)
    text(s, Inches(3.85), Inches(5.78), Inches(7.4), Inches(0.4), "Écrivez votre message…", 13, mut)
    oval(s, Inches(11.85), Inches(5.7), Inches(0.52), Inches(0.52), ACCENT)
    text(s, Inches(11.85), Inches(5.78), Inches(0.52), Inches(0.4), "→", 14, WHITE, True, PP_ALIGN.CENTER)
    footer(s, n, total)


def archives(prs, n, total):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    card, ink, mut, line = chrome(s, "Archives", "Archives", "Le titre généré remplace le libellé générique.")
    round_rect(s, Inches(2.95), Inches(1.28), Inches(2.55), Inches(0.38), PAPER, line, 0.4)
    text(s, Inches(2.95), Inches(1.32), Inches(1.25), Inches(0.3), "Textes", 11, ink, True, PP_ALIGN.CENTER)
    text(s, Inches(4.2), Inches(1.32), Inches(1.25), Inches(0.3), "Conversations", 11, mut, False, PP_ALIGN.CENTER)
    round_rect(s, Inches(5.7), Inches(1.28), Inches(6.95), Inches(0.38), PAPER, line, 0.18)
    text(s, Inches(5.9), Inches(1.32), Inches(6.5), Inches(0.3), "Rechercher", 12, mut)
    rows = [
        ("Accueil du dimanche — persévérance", "25 oct. 2026  ·  07:02", "Frères et sœurs, ce matin nous parlons de la persévérance…"),
        ("Réunion d’équipe — plan de la semaine", "24 oct. 2026  ·  18:40", "Trois priorités : le culte, les visites, le budget."),
    ]
    for i, (title, meta, excerpt) in enumerate(rows):
        y = Inches(1.85) + Inches(i * 2.25)
        round_rect(s, Inches(2.95), y, Inches(9.7), Inches(2.05), PAPER, line, 0.06)
        text(s, Inches(3.2), y + Inches(0.18), Inches(6.8), Inches(0.35), title, 16, ink, False, font=SERIF)
        text(s, Inches(9.3), y + Inches(0.22), Inches(3.1), Inches(0.28), meta, 11, mut, False, PP_ALIGN.RIGHT)
        text(s, Inches(3.2), y + Inches(0.62), Inches(9.2), Inches(0.5), excerpt, 13, mut)
        for j, a in enumerate(["Copier", "WhatsApp", "PDF", "Supprimer"]):
            x = Inches(3.2) + Inches(j * 1.7)
            round_rect(s, x, y + Inches(1.35), Inches(1.5), Inches(0.34), WHITE, line, 0.25)
            text(s, x, y + Inches(1.38), Inches(1.5), Inches(0.28), a, 11, ink, False, PP_ALIGN.CENTER)
    footer(s, n, total)


def reglages(prs, n, total):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    card, ink, mut, line = chrome(s, "Réglages", "Réglages", "Le serveur se configure ici, plus sur l’écran de connexion.")
    blocks = [
        ("Apparence", "Blanc ou noir. Le même choix est disponible en haut, face au logo."),
        ("Serveur", "Ce PC ou En ligne. Adresse Render pour le téléphone et les amis."),
        ("Compte", "Identité du compte. Déconnexion propre, session expirée gérée."),
        ("Archives", "Joindre l’audio. Copie optionnelle vers Google Drive."),
        ("Notifications", "Prévenir lorsque une longue transcription est prête."),
        ("Mobile", "Sur l’app native, le serveur distant est utilisé par défaut."),
    ]
    for i, (title, body) in enumerate(blocks):
        col, row = i % 3, i // 3
        x = Inches(2.95) + Inches(col * 3.3)
        y = Inches(1.32) + Inches(row * 2.55)
        round_rect(s, x, y, Inches(3.1), Inches(2.35), PAPER, line, 0.08)
        text(s, x + Inches(0.2), y + Inches(0.22), Inches(2.7), Inches(0.4), title, 16, ink, False, font=SERIF)
        text(s, x + Inches(0.2), y + Inches(0.75), Inches(2.7), Inches(1.3), body, 13, mut)
    footer(s, n, total)


def sombre(prs, n, total):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(s, "Transcription", "Transcription", "Fond noir. Contraste élevé. Même structure.", dark=True)
    round_rect(s, Inches(2.95), Inches(1.28), Inches(3.35), Inches(5.35), DARK_CARD, DARK_LINE, 0.06)
    round_rect(s, Inches(3.25), Inches(2.5), Inches(2.75), Inches(0.42), ACCENT, None, 0.2)
    text(s, Inches(3.25), Inches(2.56), Inches(2.75), Inches(0.32), "Commencer", 13, WHITE, True, PP_ALIGN.CENTER)
    round_rect(s, Inches(6.5), Inches(1.28), Inches(6.15), Inches(5.35), DARK_CARD, DARK_LINE, 0.06)
    text(s, Inches(6.7), Inches(3.3), Inches(5.75), Inches(0.4), "Rien pour l’instant.", 18, WHITE, False, PP_ALIGN.CENTER, SERIF)
    text(s, Inches(6.9), Inches(3.85), Inches(5.35), Inches(0.7), "Le thème se change en haut à droite,\njamais à côté du titre de page.", 13, RGBColor(0xA0, 0xA0, 0xA0), False, PP_ALIGN.CENTER)
    footer(s, n, total, dark=True)


def mobile(prs, n, total):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, PAPER)
    kicker(s, "08  ·  Natif")
    heading(s, "Android et iOS", "Même produit. Barre du bas. Serveur en ligne par défaut.")

    def phone(x, title, lines, dark=False):
        bg = DARK if dark else WHITE
        ink = WHITE if dark else INK
        mut = RGBColor(0xA0, 0xA0, 0xA0) if dark else MUTED
        stroke = DARK_LINE if dark else LINE
        round_rect(s, x, Inches(1.75), Inches(3.85), Inches(5.05), bg, stroke, 0.08)
        text(s, x + Inches(0.25), Inches(1.95), Inches(2.4), Inches(0.3), "Vocalia", 14, ink, False, font=SERIF)
        oval(s, x + Inches(3.2), Inches(1.95), Inches(0.34), Inches(0.34), DARK_CARD if dark else OFF, stroke)
        text(s, x + Inches(0.25), Inches(2.4), Inches(3.35), Inches(0.4), title, 18, ink, False, font=SERIF)
        for i, line in enumerate(lines):
            round_rect(s, x + Inches(0.25), Inches(2.95) + Inches(i * 0.72), Inches(3.35), Inches(0.58), DARK_CARD if dark else PAPER, stroke, 0.12)
            text(s, x + Inches(0.4), Inches(3.07) + Inches(i * 0.72), Inches(3.05), Inches(0.38), line, 12, ink)
        for i, name in enumerate(["Transcr.", "Assistant", "Archives", "Réglages"]):
            text(s, x + Inches(0.08) + Inches(i * 0.94), Inches(6.28), Inches(0.94), Inches(0.28), name, 9, ACCENT if i == 0 else mut, True, PP_ALIGN.CENTER)

    phone(Inches(0.55), "Transcription", ["Micro — Commencer", "Fichier audio ou vidéo", "Lien YouTube"])
    phone(Inches(4.75), "Assistant", ["Message écrit", "Question vocale", "Archiver l’échange"])
    phone(Inches(8.95), "Archives", ["Titre automatique", "WhatsApp et PDF", "Suppression"], dark=True)
    footer(s, n, total)


def partage(prs, n, total):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, PAPER)
    kicker(s, "Fonctions")
    heading(s, "Titre et diffusion")
    items = [
        ("Titre automatique", "Quatre à dix mots en français, générés à l’archivage, éditables."),
        ("Copier", "Titre et texte dans le presse-papiers, prêts à coller."),
        ("WhatsApp", "Ouverture directe avec le contenu, téléphone ou ordinateur."),
        ("PDF", "Mise en page simple, impression ou enregistrement PDF."),
        ("Partage système", "Feuille de partage du téléphone lorsqu’elle existe."),
        ("Google Drive", "Option, dans Réglages, après liaison du compte."),
    ]
    for i, (title, body) in enumerate(items):
        col, row = i % 3, i // 3
        x = Inches(0.5) + Inches(col * 4.2)
        y = Inches(1.75) + Inches(row * 2.45)
        round_rect(s, x, y, Inches(3.95), Inches(2.2), WHITE, LINE, 0.08)
        rect(s, x, y, Inches(0.09), Inches(2.2), ACCENT)
        text(s, x + Inches(0.35), y + Inches(0.28), Inches(3.35), Inches(0.45), title, 16, INK, False, font=SERIF)
        text(s, x + Inches(0.35), y + Inches(0.85), Inches(3.35), Inches(1.05), body, 13, MUTED)
    footer(s, n, total)


def cloture(prs, n, total):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    rect(s, 0, 0, W, H, INK)
    text(s, Inches(0.7), Inches(1.7), Inches(11.8), Inches(0.3), "PROCHAINES ÉTAPES", 12, GOLD, True)
    text(s, Inches(0.7), Inches(2.15), Inches(12), Inches(0.8), "Tester, puis publier.", 36, WHITE, False, font=SERIF)
    text(
        s,
        Inches(0.7),
        Inches(3.2),
        Inches(11.5),
        Inches(2.4),
        "Web\nhttps://vocalia-ud4o.onrender.com\n\n"
        "Android  —  npm run mobile   puis   npm run android:open\n"
        "iOS  —  mêmes commandes, sur un Mac, avec Xcode",
        16,
        RGBColor(0xC4, 0xC4, 0xC4),
    )
    text(s, Inches(0.7), Inches(6.85), Inches(11), Inches(0.3), "Le premier chargement web peut afficher l’écran Render, le temps de réveiller le service.", 12, RGBColor(0x8A, 0x8A, 0x8A))
    footer(s, n, total, dark=True)


def main():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    builders = [
        cover,
        sommaire,
        parcours,
        splash,
        lambda p, i, t: auth(p, i, t, False),
        lambda p, i, t: auth(p, i, t, True),
        studio,
        document,
        assistant,
        archives,
        reglages,
        sombre,
        mobile,
        partage,
        cloture,
    ]
    total = len(builders)
    for i, fn in enumerate(builders, 1):
        fn(prs, i, total)
    out = Path(__file__).resolve().parent / "Vocalia-maquette-produit.pptx"
    prs.save(out)
    print(out)


if __name__ == "__main__":
    main()
